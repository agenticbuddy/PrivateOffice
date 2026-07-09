"""Blank OOXML document generation.

New documents are created from the library's neutral OOXML defaults. We deliberately do
NOT stamp a content/default language onto the document: the Collabora/LibreOffice core
derives metric-field formatting (e.g. the point measurement unit) and the status-bar
language tag from the document's CONTENT language, not from the editor UI language. Pinning
that content language to the creator's UI locale therefore made the unit render in that
locale (e.g. Cyrillic "пт") for every viewer regardless of their UI language. The editor's
UI language is supplied separately at open time via the ``lang`` URL parameter, and the
creator's locale is preserved independently in the ``creator_locale`` node metadata.
"""
from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pptx import Presentation


def _docx() -> bytes:
    doc = Document()
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# LibreOffice's localized default sheet-name prefix, by language. The first sheet of a NEW
# workbook is named "<prefix>1" so it matches the editor's own localized new sheets ("Лист2", …)
# instead of openpyxl's hard-coded English "Sheet". Only the sheet NAME is localized — plain
# content shared by every viewer; we still do NOT stamp a content language (that would change
# metric-unit rendering for all viewers, see the module docstring).
_SHEET_PREFIX = {
    "en": "Sheet", "ru": "Лист", "uk": "Аркуш", "de": "Tabelle", "fr": "Feuille",
    "es": "Hoja", "it": "Foglio", "pt": "Folha", "pt-br": "Planilha", "nl": "Blad",
    "pl": "Arkusz", "cs": "List", "sk": "Hárok", "tr": "Sayfa", "sv": "Blad",
    "fi": "Taulukko", "da": "Ark", "nb": "Ark", "hu": "Munkalap", "ro": "Foaie",
    "bg": "Лист", "el": "Φύλλο", "ca": "Full", "ja": "シート", "ko": "시트",
    "zh-cn": "工作表", "zh-tw": "工作表", "ar": "ورقة", "he": "גיליון",
}


def _sheet_title(locale: str | None) -> str:
    """Localized first-sheet title, e.g. ru → "Лист1", en → "Sheet1" (fallback "Sheet1")."""
    if not locale:
        return "Sheet1"
    key = locale.strip().lower().replace("_", "-")
    prefix = _SHEET_PREFIX.get(key) or _SHEET_PREFIX.get(key.split("-")[0]) or "Sheet"
    return f"{prefix}1"


def _xlsx(locale: str | None = None) -> bytes:
    wb = Workbook()
    wb.active.title = _sheet_title(locale)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    prs = Presentation()
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


_BUILDERS = {"docx": _docx, "xlsx": _xlsx, "pptx": _pptx}


def blank_document(fmt: str, locale: str | None = None) -> bytes:
    """Return bytes of a blank document of the given format (neutral content language).

    ``locale`` only localizes the xlsx first-sheet NAME (see ``_sheet_title``); docx/pptx ignore it.
    """
    if fmt not in _BUILDERS:
        raise ValueError(f"Unsupported creatable format: {fmt}")
    return _xlsx(locale) if fmt == "xlsx" else _BUILDERS[fmt]()
